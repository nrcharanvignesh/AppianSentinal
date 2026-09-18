"""Lazy access to the shipped GenAI proxy workbook."""

from __future__ import annotations

import logging
import re
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path
from types import MappingProxyType
from typing import Literal, Mapping
from urllib.parse import urlsplit

from openpyxl import load_workbook

from appian_sentinel.knowledge import resolve_ground_truth_path

_LOGGER = logging.getLogger(__name__)
_WORKBOOK_NAME = "GenAI Documentation.xlsx"
_URL_RE = re.compile(r"https?://[^\s\"']+")

ModelProtocol = Literal["openai", "anthropic"]


@dataclass(frozen=True)
class GenAIModel:
    """One model documented by the GenAI proxy workbook."""

    name: str
    canonical_name: str
    aliases: tuple[str, ...]
    provider: str
    model_type: str
    protocol: ModelProtocol
    region: str | None
    max_input_tokens: int | None
    max_output_tokens: int | None
    deployment_count: int | None
    requests_per_minute: int | None = None
    tokens_per_minute: int | None = None


@dataclass(frozen=True)
class GenAICatalog:
    """Structured proxy endpoints and models from the workbook."""

    proxy_base_urls: tuple[str, ...]
    models: Mapping[str, GenAIModel]


@dataclass(frozen=True)
class GenAIConfigDefaults:
    """Workbook-backed defaults consumed by application settings."""

    base_url: str
    primary_model: str
    fast_model: str


def _as_text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _as_optional_int(value: object) -> int | None:
    if value is None or value == "":
        return None
    return int(value)


def _protocol(name: str, canonical_name: str, provider: str) -> ModelProtocol:
    normalized = f"{name}.{canonical_name}.{provider}".casefold().replace("/", ".")
    return "anthropic" if ".anthropic." in f".{normalized}." else "openai"


def _proxy_base_urls(endpoint_descriptions: list[str]) -> tuple[str, ...]:
    bases: set[str] = set()
    for description in endpoint_descriptions:
        for raw_url in _URL_RE.findall(description):
            parsed = urlsplit(raw_url.rstrip(".,);"))
            if parsed.hostname not in {"localhost", "127.0.0.1"}:
                continue
            bases.add(f"{parsed.scheme}://{parsed.netloc}")
    return tuple(sorted(bases))


def _parse_models(rows: list[tuple[object, ...]]) -> Mapping[str, GenAIModel]:
    if not rows:
        raise ValueError("Models List is empty")
    headers = {_as_text(value): index for index, value in enumerate(rows[0])}
    required = {
        "Model Name",
        "Canonical Model",
        "Provider",
        "Model Type",
        "Max Input Tokens",
        "Max Output Tokens",
    }
    if not required.issubset(headers):
        raise ValueError("Models List does not contain the expected columns")

    def cell(row: tuple[object, ...], header: str) -> object:
        index = headers[header]
        return row[index] if index < len(row) else None

    models: dict[str, GenAIModel] = {}
    for row in rows[1:]:
        name = _as_text(cell(row, "Model Name"))
        if not name:
            continue
        canonical_name = _as_text(cell(row, "Canonical Model"))
        provider = _as_text(cell(row, "Provider"))
        aliases = tuple(
            alias.strip()
            for alias in _as_text(cell(row, "Aliases")).split(";")
            if alias.strip()
        )
        model = GenAIModel(
            name=name,
            canonical_name=canonical_name,
            aliases=aliases,
            provider=provider,
            model_type=_as_text(cell(row, "Model Type")),
            protocol=_protocol(name, canonical_name, provider),
            region=_as_text(cell(row, "Region")) or None,
            max_input_tokens=_as_optional_int(cell(row, "Max Input Tokens")),
            max_output_tokens=_as_optional_int(cell(row, "Max Output Tokens")),
            deployment_count=_as_optional_int(cell(row, "Deployment Count")),
        )
        models[name.casefold()] = model
    if not models:
        raise ValueError("Models List did not contain any models")
    return MappingProxyType(models)


@lru_cache(maxsize=None)
def _load_catalog(path: Path) -> GenAICatalog | None:
    if not path.is_file():
        _LOGGER.warning(
            "[WARN] %s is missing; using built-in GenAI configuration defaults",
            path,
        )
        return None
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            model_rows = list(workbook["Models List"].iter_rows(values_only=True))
            endpoint_rows = list(workbook["Endpoints"].iter_rows(values_only=True))
        finally:
            workbook.close()
        descriptions = [
            _as_text(value)
            for row in endpoint_rows
            for value in row
            if value is not None
        ]
        return GenAICatalog(
            proxy_base_urls=_proxy_base_urls(descriptions),
            models=_parse_models(model_rows),
        )
    except (KeyError, OSError, TypeError, ValueError) as exc:
        _LOGGER.warning(
            "[WARN] Could not load %s (%s); using built-in GenAI configuration defaults",
            path,
            exc,
        )
        return None


def load_genai_catalog() -> GenAICatalog | None:
    """Load and cache the workbook, or return ``None`` on fallback."""
    return _load_catalog(resolve_ground_truth_path(_WORKBOOK_NAME))


def get_models() -> Mapping[str, GenAIModel] | None:
    """Return models keyed by normalized model name."""
    catalog = load_genai_catalog()
    return None if catalog is None else catalog.models


def get_model(name: str) -> GenAIModel | None:
    """Return documented metadata and protocol for one model."""
    models = get_models()
    return None if models is None else models.get(name.casefold())


def get_proxy_base_urls() -> tuple[str, ...] | None:
    """Return proxy base URLs documented in workbook request examples."""
    catalog = load_genai_catalog()
    return None if catalog is None else catalog.proxy_base_urls


def get_config_defaults(
    fallback_base_url: str,
    fallback_primary_model: str,
    fallback_fast_model: str,
) -> GenAIConfigDefaults | None:
    """Return settings defaults validated against the workbook catalog."""
    catalog = load_genai_catalog()
    if catalog is None:
        return None
    chat_models = tuple(
        model for model in catalog.models.values() if model.model_type == "chat"
    )
    if not chat_models:
        return None

    def select(preferred: str, family: str) -> str:
        documented = catalog.models.get(preferred.casefold())
        if documented is not None:
            return documented.name
        matches = sorted(
            model.name
            for model in chat_models
            if model.protocol == "anthropic" and family in model.name.casefold()
        )
        return matches[-1] if matches else chat_models[0].name

    base_url = (
        catalog.proxy_base_urls[0]
        if catalog.proxy_base_urls
        else fallback_base_url
    )
    return GenAIConfigDefaults(
        base_url=base_url,
        primary_model=select(fallback_primary_model, "opus"),
        fast_model=select(fallback_fast_model, "sonnet"),
    )


def clear_genai_cache() -> None:
    """Clear cached workbook data for tests or changed source files."""
    _load_catalog.cache_clear()
